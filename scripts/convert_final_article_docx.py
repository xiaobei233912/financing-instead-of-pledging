"""Literal source-to-Word conversion; no editorial changes to the source text."""
from pathlib import Path
import re,json,hashlib,zipfile
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
SOURCE=ROOT/'output/融资代替质押详解.md'
DEST=SOURCE.with_suffix('.docx')
QA=ROOT/'tmp/final_article_docx_qa';QA.mkdir(parents=True,exist_ok=True)
raw=SOURCE.read_bytes();source_hash=hashlib.sha256(raw).hexdigest()
lines=raw.decode('utf-8-sig').splitlines()
doc=Document();sec=doc.sections[0]
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
sec.footer_distance=Inches(.32)

def font(obj,name='SimSun',size=11,bold=False):
    obj.font.name=name;obj.font.size=Pt(size);obj.font.bold=bold;obj.font.color.rgb=RGBColor(0,0,0)
    rp=obj.element.get_or_add_rPr()
    fonts=rp.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rp.insert(0,fonts)
    for k in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+k),name)
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+k),None)

for name,size,bold,face in [('Normal',11,False,'SimSun'),('Title',21,True,'Microsoft YaHei'),
                             ('Heading 1',15,True,'Microsoft YaHei'),('Heading 2',12,True,'Microsoft YaHei'),
                             ('Caption',9.5,False,'Microsoft YaHei')]:
    st=doc.styles[name];font(st,face,size,bold)
    pf=st.paragraph_format;pf.line_spacing=1.3;pf.space_after=Pt(6);pf.widow_control=True
    if name.startswith('Heading'):pf.keep_with_next=True;pf.space_before=Pt(14);pf.space_after=Pt(7)
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Title'].paragraph_format.keep_with_next=True

expected=[];images=[];table_count=0
def para(text,style=None):
    p=doc.add_paragraph(style=style)
    # Emphasize original rule labels without changing a single character.
    m=re.match(r'^(规则[一二三四]，)',text) if style is None else None
    if m:
        p.add_run(m.group(1)).bold=True;p.add_run(text[len(m.group(1)):])
    else:p.add_run(text)
    expected.append(text)
    return p

def table_line(line):return line.strip().startswith('|') and line.strip().endswith('|')
def cells(line):return [x.strip() for x in line.strip()[1:-1].split('|')]
def separator(line):return table_line(line) and all(re.fullmatch(r':?-{3,}:?',x) for x in cells(line))

i=0
while i<len(lines):
    s=lines[i]
    if not s.strip():i+=1;continue
    if table_line(s) and i+1<len(lines) and separator(lines[i+1]):
        rows=[cells(s)];i+=2
        while i<len(lines) and table_line(lines[i]):rows.append(cells(lines[i]));i+=1
        assert all(len(row)==len(rows[0]) for row in rows)
        table=doc.add_table(rows=0,cols=len(rows[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        widths=[2.28,1.54,1.54,1.54]
        for col,width in zip(table.columns,widths):col.width=Inches(width)
        pr=table._tbl.tblPr;borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        pr.append(borders)
        for ri,rowdata in enumerate(rows):
            row=table.add_row();trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if ri==0:
                header=OxmlElement('w:tblHeader');header.set(qn('w:val'),'true');trpr.append(header)
            for ci,(cell,value) in enumerate(zip(row.cells,rowdata)):
                cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                for side,v in [('top',85),('bottom',85),('left',105),('right',105)]:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');marg.append(e)
                tcpr.append(marg)
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'DCE5EA' if ri==0 else ('F6F8FA' if ri%2==0 else 'FFFFFF'));tcpr.append(shd)
                p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.16;p.paragraph_format.keep_with_next=ri==0
                run=p.add_run(value);font(run,'Microsoft YaHei',9.5,ri==0);expected.append(value)
        table_count+=1
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
        continue
    m=re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)',s)
    if m:
        alt,reference=m.groups();candidates=[SOURCE.parent/reference,ROOT/'reports'/reference]
        path=next((p for p in candidates if p.exists()),None)
        if path is None:raise FileNotFoundError(reference)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(3)
        shape=p.add_run().add_picture(str(path),width=Inches(6.75));shape._inline.docPr.set('descr',alt);shape._inline.docPr.set('title',reference)
        cap=para(alt,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.space_after=Pt(10)
        images.append(dict(reference=reference,resolved=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    elif i==0:
        para(s,'Title')
    elif s.startswith('作者：'):
        p=para(s);p.paragraph_format.space_after=Pt(12)
        for run in p.runs:font(run,'Microsoft YaHei',10)
    elif re.match(r'^[一二三四五六七八九十]+[、.]',s):
        para(s,'Heading 1')
    elif s in ('1.再平衡时的调仓灵活性问题。','2.担保资产转出控制线300%导致没破产但没钱花的问题。'):
        para(s,'Heading 2')
    else:para(s)
    i+=1

# Page numbers are layout fields, outside the original article body.
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run();font(run,'Calibri',9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title=lines[0];doc.core_properties.author=lines[1].removeprefix('作者：')
doc.core_properties.subject='';doc.core_properties.comments=''
doc.save(DEST)

with zipfile.ZipFile(DEST) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    actual=[''.join(p.itertext()) for p in []]
    actual=[''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in xml.xpath('.//w:body//w:p',namespaces=ns)]
    actual=[x for x in actual if x]
    assert actual==expected,[(j,a,b) for j,(a,b) in enumerate(zip(actual,expected)) if a!=b][:3]
    media=[hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')]
    assert sorted(media)==sorted(x['sha256'] for x in images)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
verification=dict(source=str(SOURCE),output=str(DEST),source_sha256=source_hash,source_unchanged=True,
                  text_blocks=len(expected),exact_text_match=True,tables=table_count,images=images,
                  conversion='Markdown table delimiters become native tables; image alt text retained as exact captions; all prose and cell text compared in order.')
(QA/'text_verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in verification.items() if k!='images'},ensure_ascii=True,indent=2))
